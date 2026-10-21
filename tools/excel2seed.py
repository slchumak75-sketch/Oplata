#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Конвертер исходного Excel-файла «Оплаты Строймаркет.xlsx» в seed.json приложения Oplata.

Зачем: при первом запуске приложение загружает pkg/store/seed.json, чтобы пользователь
сразу видел свои данные за 2025 год без ручного импорта.

Использование:
    pip install openpyxl
    python3 tools/excel2seed.py "Оплаты Строймаркет.xlsx" pkg/store/seed.json

Правила преобразования (подробно — docs/EXCEL_FORMAT.md):
  * Лист называется годом («2025»). Строка 1 — месяц, строка 2 — день, строка 3 — тип колонки.
  * Колонки идут парами на каждую дату: «Собств» (сумма к оплате за поставленный товар,
    дата — срок оплаты с учётом отсрочки) и «Выплаты» (сумма выплаты, в Excel хранится со знаком минус).
  * «Собств»  -> obligations: amount_cents = round(value * 100) (положительное начисление долга).
  * «Выплаты» -> payments:    amount_cents = round(-value * 100) (выплата уменьшает долг;
    положительное значение в Excel — ошибка знака, импортируется как отрицательная выплата + предупреждение).
  * Поставщики — столбец A ниже строки «ПОСТАВЩИК» (строки 6..48), порядок сохраняется.
"""
import json
import sys
from collections import OrderedDict

import openpyxl

MONTHS = {
    "январь": 1, "февраль": 2, "март": 3, "апрель": 4, "май": 5,
    "июнь": 6, "июню": 6,  # в исходном файле опечатка «июню»
    "июль": 7, "август": 8, "сентябрь": 9, "октябрь": 10, "ноябрь": 11, "декабрь": 12,
}


def to_cents(v: float) -> int:
    return int(round(v * 100))


def main() -> int:
    src = sys.argv[1] if len(sys.argv) > 1 else "Оплаты Строймаркет.xlsx"
    dst = sys.argv[2] if len(sys.argv) > 2 else "pkg/store/seed.json"

    wb = openpyxl.load_workbook(src, data_only=True, read_only=True)
    ws = wb[wb.sheetnames[0]]
    sheet = ws.title.strip()
    year = int(sheet) if sheet.isdigit() else 2025

    rows = list(ws.iter_rows(min_row=1, max_row=200, max_col=2000, values_only=True))
    r_month, r_day, r_kind = rows[0], rows[1], rows[2]

    # Колонки C.. : восстанавливаем (дата, тип). Месяц/день в колонках «Выплаты» пустые — берём от пары.
    columns = []  # (col_index_0based, iso_date, kind) kind: 'sobstv' | 'vyplaty'
    cur_month = None
    cur_day = None
    for idx in range(2, len(r_kind)):
        kind_raw = (str(r_kind[idx]).strip() if r_kind[idx] is not None else "")
        if r_month[idx] is not None and str(r_month[idx]).strip():
            cur_month = MONTHS.get(str(r_month[idx]).strip().lower())
        if r_day[idx] is not None:
            cur_day = int(r_day[idx])
        if kind_raw == "Собств":
            kind = "sobstv"
        elif kind_raw == "Выплаты":
            kind = "vyplaty"
        else:
            continue
        if cur_month is None or cur_day is None:
            raise SystemExit(f"Колонка {idx + 1}: нет месяца/дня")
        columns.append((idx, f"{year:04d}-{cur_month:02d}-{cur_day:02d}", kind))

    # Поставщики: строки ниже заголовка «ПОСТАВЩИК» (строка 5 => индекс 4).
    header_idx = None
    for i, row in enumerate(rows[:10]):
        if row and str(row[0] or "").strip() == "ПОСТАВЩИК":
            header_idx = i
            break
    if header_idx is None:
        raise SystemExit("Не найдена строка заголовка ПОСТАВЩИК")

    suppliers, obligations, payments, warnings = [], [], [], []
    per_supplier_balance = OrderedDict()
    obl_n = pay_n = 0
    for i in range(header_idx + 1, len(rows)):
        row = rows[i]
        name = str(row[0]).strip() if row and row[0] is not None else ""
        if not name:
            break  # данные идут сплошным блоком
        sid = f"sup-{len(suppliers) + 1:03d}"
        suppliers.append(OrderedDict([
            ("id", sid), ("name", name), ("unp", ""), ("deferral_days", 0), ("note", ""), ("archived", False),
        ]))
        balance = 0
        for col_idx, iso, kind in columns:
            v = row[col_idx] if col_idx < len(row) else None
            if v is None or not isinstance(v, (int, float)) or v == 0:
                continue
            if kind == "sobstv":
                obl_n += 1
                cents = to_cents(float(v))
                balance += cents
                obligations.append(OrderedDict([
                    ("id", f"obl-{obl_n:06d}"), ("supplier_id", sid), ("due_date", iso),
                    ("amount_cents", cents), ("delivery_date", ""), ("doc_number", ""), ("comment", ""),
                ]))
            else:
                pay_n += 1
                cents = to_cents(-float(v))  # Excel: выплата -1013.49 -> +101349 коп. выплачено
                balance -= cents
                comment = ""
                if v > 0:
                    warnings.append(
                        f"{name}, {iso}: «Выплаты» с положительным знаком ({v}) — в Excel выплаты "
                        f"принято писать с минусом. Импортировано как отрицательная выплата "
                        f"(увеличивает задолженность), проверьте запись."
                    )
                    comment = "Импорт из Excel: выплата была с положительным знаком, проверьте."
                payments.append(OrderedDict([
                    ("id", f"pay-{pay_n:06d}"), ("supplier_id", sid), ("date", iso),
                    ("amount_cents", cents), ("doc_number", ""), ("comment", comment),
                ]))
        per_supplier_balance[name] = balance

    data = OrderedDict([
        ("version", 3),  # схема v3: v2 (unp/buyers) + необязательная разбивка начислений по НДС
        ("settings", OrderedDict([
            ("company_name", "Строймаркет"),
            ("currency_label", "руб."),
            ("warn_days", 3),
            ("seeded_from", f"{src} (лист {sheet})"),
        ])),
        ("suppliers", suppliers),
        ("obligations", obligations),
        ("payments", payments),
        ("buyers", []),  # в Excel покупателей нет — справочник заполняется в приложении
        ("warnings", warnings),
    ])
    with open(dst, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=1)
        f.write("\n")

    total_obl = sum(o["amount_cents"] for o in obligations)
    total_pay = sum(p["amount_cents"] for p in payments)
    print(f"OK -> {dst}")
    print(f"  поставщиков: {len(suppliers)}  начислений(Собств): {len(obligations)}  выплат: {len(payments)}")
    print(f"  сумма Собств: {total_obl / 100:,.2f}  сумма выплат: {total_pay / 100:,.2f}  ЗАДОЛЖЕННОСТЬ: {(total_obl - total_pay) / 100:,.2f}")
    print("  задолженность по поставщикам (ненулевая):")
    for n, b in per_supplier_balance.items():
        if b:
            print(f"    {n}: {b / 100:,.2f}")
    for w in warnings:
        print(f"  ПРЕДУПРЕЖДЕНИЕ: {w}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
