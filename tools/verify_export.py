#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Независимая проверка XLSX, который генерирует приложение (pkg/excelio):
открываем файл сторонней библиотекой openpyxl и сверяем структуру и итоги
с исходным «Оплаты Строймаркет.xlsx» (docs/EXCEL_FORMAT.md).

Использование (CI делает именно так):
    go run ./tools/exportdemo /tmp/oplata-2025.xlsx
    python3 tools/verify_export.py /tmp/oplata-2025.xlsx
"""
import sys

import openpyxl

EXPECTED_DEBTORS = {
    "АВИДО": 1590.45,
    "Алваис": 234.76,
    "Продмир": 1735.68,
    "Культбытхозторг": 856.41,
    "ВергоХоз": 485.51,
    "Система ООО": 16491.64,
    "ТД Комплект": 4360.47,
}
EXPECTED_TOTAL = 25754.92


def main() -> int:
    path = sys.argv[1] if len(sys.argv) > 1 else "oplata-2025.xlsx"
    wb = openpyxl.load_workbook(path, data_only=True)
    assert wb.sheetnames == ["2025", "Операции"], f"листы: {wb.sheetnames}"
    ws = wb["2025"]
    assert ws["A4"].value == "ВСЕГО", ws["A4"].value
    assert ws["A5"].value == "ПОСТАВЩИК", ws["A5"].value
    # В оригинале написано «ЗАДОЛЖЕНОСТЬ» с одной Н; экспорт пишет орфографически правильно.
    assert ws["B3"].value == "ЗАДОЛЖЕННОСТЬ", ws["B3"].value
    assert str(ws.freeze_panes) == "C5", ws.freeze_panes

    # Шапка дат: пары Собств/Выплаты от C до ABD (365 дат), дни 1..31 января.
    assert ws["C1"].value == "январь" and ws["C2"].value == 1
    assert ws["C3"].value == "Собств" and ws["D3"].value == "Выплаты"
    assert ws["ABC1"].value == "декабрь" and ws["ABC2"].value == 31

    # Кэшированные значения формул столбца B должны совпадать с Excel-оригиналом.
    debts = {}
    row = 6
    while ws.cell(row=row, column=1).value:
        name = ws.cell(row=row, column=1).value
        debts[name] = round(float(ws.cell(row=row, column=2).value or 0), 2)
        row += 1
    assert len(debts) == 43, f"поставщиков: {len(debts)}"
    nonzero = {k: v for k, v in debts.items() if v != 0}
    assert nonzero == EXPECTED_DEBTORS, f"должники: {nonzero}"
    total_b4 = round(float(ws["B4"].value), 2)
    assert total_b4 == EXPECTED_TOTAL, f"B4={total_b4}"

    # Прямой пересчёт ячеек (без формул): Собств положительны, Выплаты отрицательны —
    # кроме одной известной аномалии оригинала (Алваис, +117,38; экспорт её воспроизводит).
    sobstv_sum = vyplaty_sum = 0.0
    positive_vyplaty = []
    for r in range(6, 6 + 43):
        for c in range(3, 732 + 1):
            v = ws.cell(row=r, column=c).value
            if isinstance(v, (int, float)) and v:
                if ws.cell(row=3, column=c).value == "Собств":
                    assert v > 0, (r, c, v)
                    sobstv_sum += v
                else:
                    if v > 0:
                        positive_vyplaty.append((ws.cell(row=r, column=1).value, v))
                    vyplaty_sum += v
    assert positive_vyplaty == [("Алваис", 117.38)], positive_vyplaty
    assert round(sobstv_sum, 2) == 837334.34, round(sobstv_sum, 2)
    assert round(sobstv_sum + vyplaty_sum, 2) == EXPECTED_TOTAL

    # Лист «Операции»: заголовок + 853 операции исходного Excel.
    ops = wb["Операции"]
    assert ops["A1"].value == "Тип" and ops["B1"].value == "Поставщик"
    assert ops.max_row == 1 + 438 + 415, f"операций строк: {ops.max_row}"

    print(f"OK: {path} — структура как в оригинале, итоги сошлись "
          f"(Собств {sobstv_sum:,.2f}, долг {total_b4:,.2f}, поставщиков 43)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
